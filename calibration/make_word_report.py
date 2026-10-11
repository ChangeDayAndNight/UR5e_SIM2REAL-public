# -*- coding: utf-8 -*-
"""Generate a detailed Word (.docx) report of the hand-eye calibration project, incl. math."""
import json
import numpy as np
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

R = json.load(open("captures_final/results/calibration_results.json"))
D = json.load(open("captures_final/results/demo_results.json"))
T = np.array(R["T_EE_CAM"]); Tbb = np.array(R["T_base_board"]); intr = R["intrinsics"]; err = R["error"]
NAVY = RGBColor(0x1F, 0x3A, 0x5F)

doc = Document()
st = doc.styles["Normal"]; st.font.name = "맑은 고딕"; st.font.size = Pt(10.5)

def h1(t):
    p = doc.add_heading(t, level=1)
    for r in p.runs: r.font.color.rgb = NAVY
    return p
def h2(t):
    p = doc.add_heading(t, level=2)
    for r in p.runs: r.font.color.rgb = NAVY
    return p
def para(t, bold=False, size=10.5):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = bold; r.font.size = Pt(size); return p
def bullet(t, lvl=0):
    p = doc.add_paragraph(style="List Bullet" if lvl==0 else "List Bullet 2"); p.add_run(t); return p
def mono(t):
    p = doc.add_paragraph()
    for line in t.split("\n"):
        r = p.add_run(line + "\n"); r.font.name = "Consolas"; r.font.size = Pt(9.5)
    return p
def mat(M, name=""):
    s = (name + " =\n" if name else "")
    s += "\n".join("[ " + "  ".join(f"{v:10.6f}" for v in row) + " ]" for row in M)
    mono(s)

# ================= TITLE =================
tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tp.add_run("Eye-in-Hand 핸드-아이 캘리브레이션 기술 보고서"); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = NAVY
sp = doc.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sp.add_run("RealSense D435i on UR5e · ChArUco + 번들조정 · RTDE\n수학적 원리 · 과거 측정 데이터 · 최종 결과 · 오프라인 변환")
sr.font.size = Pt(11); sr.font.color.rgb = RGBColor(0x55,0x55,0x55)
doc.add_paragraph()

# ================= 0. 요약 =================
h1("0. 요약 (Executive Summary)")
para("본 프로젝트의 목표는 로봇 플랜지(손목 끝)에 강체로 장착된 카메라의 상대 자세 T_EE_CAM 을 구하는 것이다. "
     "이 행렬을 알면 카메라가 관측한 임의의 점을 로봇 베이스 좌표계로 변환할 수 있어, 도달성 분석과 정밀검사에 사용할 수 있다.")
para("최종적으로 11×8 ChArUco 보드를 27개 자세에서 촬영(모든 뷰 70/70 코너 검출)하고, 카메라 내부파라미터를 별도로 "
     "교정해 고정한 뒤(방식 A), 재투영 오차를 최소화하는 번들조정으로 T_EE_CAM 과 보드 자세를 동시에 구했다.")
para(f"결과: 카메라는 플랜지 기준 오프셋 {np.round(np.array(R['camera_offset_mm']),1).tolist()} mm, 회전 {R['camera_rot_deg']:.2f}°에 장착되어 있으며, "
     f"재투영 RMS = {err['reproj_rms_px']:.2f} px = {err['reproj_rms_mm']:.2f} mm (작업거리 {R['working_distance_m']:.3f} m). "
     "카메라→베이스 변환의 과거 측정 결과를 보존한다. 이 보고서는 현재 실기 모션 승인이 아니다.", bold=True)

# ================= 1. 시스템 =================
h1("1. 시스템 구성")
bullet("구성 방식: eye-in-hand — 카메라가 로봇 손목에 고정되어 로봇과 함께 움직이고, 보드는 작업공간에 고정.")
bullet("로봇: UR5e endpoint는 ignored configs/real.local.yaml에서 로드. 자세는 getActualTCPPose() = base→TCP(플랜지) [x,y,z,rx,ry,rz] (m, 축-각).")
bullet("카메라: eye-in-hand Intel RealSense D435i. 정확한 시리얼은 ignored configs/real.local.yaml에서 로드하며 자동 첫 장치 선택을 금지.")
bullet("보드: 11×8 ChArUco, 정사각 15 mm, 마커 11 mm, DICT_4X4_50, LEGACY 패턴.")
bullet("소프트웨어: 잠금된 Windows 배포 런타임 — pyrealsense2, ur-rtde, OpenCV, SciPy.")

# ================= 2. 수학적 원리 =================
h1("2. 수학적 원리")

h2("2.1 강체변환과 동차좌표")
para("3차원 강체변환은 회전 R(3×3, R∈SO(3), RᵀR=I, det R=+1)과 병진 t(3×1)로 이루어지며, 4×4 동차행렬로 표현한다:")
mono("T = [ R  t ]      점 변환:  [p'; 1] = T · [p; 1]   즉  p' = R·p + t\n    [ 0  1 ]")
para("합성은 행렬 곱, 역변환은  T⁻¹ = [ Rᵀ  −Rᵀt ; 0  1 ].  본 보고서의 표기:")
bullet("T_base_flange : 플랜지 좌표를 베이스 좌표로 보내는 변환 (= RTDE가 주는 TCP 자세).")
bullet("T_EE_CAM : 카메라 좌표를 플랜지 좌표로 보내는 변환 (구하려는 값). p_flange = T_EE_CAM · p_camera.")
bullet("따라서 카메라의 베이스 기준 자세:  T_base_cam = T_base_flange · T_EE_CAM.")
para("축-각(로드리게스) 표현: 회전벡터 r(3×1)의 크기 θ=‖r‖가 회전각, 방향 r/θ가 회전축. R = I + sinθ·[k]× + (1−cosθ)·[k]×²  (k=r/θ, [k]×는 반대칭행렬). OpenCV의 cv2.Rodrigues가 r↔R 변환.")

h2("2.2 카메라 모델 (핀홀 + 렌즈 왜곡)")
para("카메라 좌표계의 점 X_c=(X,Y,Z)를 정규화: x=X/Z, y=Y/Z. 렌즈 왜곡(방사 k1,k2,k3 · 접선 p1,p2)을 적용:")
mono("r² = x²+y²\n"
     "x_d = x(1 + k1 r² + k2 r⁴ + k3 r⁶) + 2p1 x y + p2(r²+2x²)\n"
     "y_d = y(1 + k1 r² + k2 r⁴ + k3 r⁶) + p1(r²+2y²) + 2p2 x y")
para("픽셀 좌표로의 투영 (내부행렬 K):")
mono("u = fx·x_d + cx\nv = fy·y_d + cy        K = [ fx 0 cx ; 0 fy cy ; 0 0 1 ]")
para("fx,fy: 초점거리(픽셀), (cx,cy): 주점. 이 (fx,fy,cx,cy,k1..k3)를 카메라 내부파라미터(intrinsics)라 한다.")

h2("2.3 ChArUco 검출과 PnP (보드 자세 추정)")
para("ChArUco는 체스보드 + ArUco 마커의 혼합 패턴이다. 마커로 각 코너의 ID를 명확히 식별하므로, 보드가 일부만 보여도 "
     "코너의 3D 보드좌표 ↔ 2D 픽셀 대응을 얻는다. (주의: OpenCV ≥4.6은 코너-마커 배치 규칙이 바뀌어, 구버전 생성 보드는 "
     "setLegacyPattern(True) 가 필요하다.)")
para("대응점 {(P_i^board, u_i)} 으로 PnP(Perspective-n-Point)를 풀어 보드→카메라 자세 (R,t)=T_cam_board 를 구한다. "
     "재투영오차를 최소화하는 문제:")
mono("min_(R,t)  Σ_i ‖ π(K, dist, R·P_i^board + t) − u_i ‖²      (π: 위 투영함수)")

h2("2.4 Hand-eye 문제: AX = XB")
para("eye-in-hand에서 두 자세 i,j 사이의 상대운동을 생각한다. 보드는 고정이므로:")
bullet("A = 플랜지의 상대운동 (베이스 기준, 로봇에서 알 수 있음) = T_base_flange,i⁻¹ · T_base_flange,j")
bullet("B = 카메라가 본 보드의 상대운동 = T_cam_board,i · T_cam_board,j⁻¹")
para("미지의 손-눈 변환 X = T_EE_CAM 는 모든 쌍에서 다음을 만족한다:")
mono("A · X = X · B")
para("회전부와 병진부로 분리하면  R_A R_X = R_X R_B  및  (R_A − I) t_X = R_X t_B − t_A. "
     "정확도의 관건: 회전축이 서로 다른(비평행) 자세가 많을수록 R_X가 잘 구속되고, 병진이 다양할수록 t_X가 잘 구속된다.")

h2("2.5 Closed-form 해 (Tsai/Park) — 초기값")
para("Tsai, Park 등은 위 식을 선형/리대수(so(3)) 형태로 풀어 X의 닫힌형 해를 준다(OpenCV cv2.calibrateHandEye). "
     "빠르지만 각 뷰의 PnP 자세만 사용하고 픽셀 잔차를 직접 최소화하지 않아 최적은 아니다. 본 프로젝트에서는 이 해를 "
     "번들조정의 초기값으로만 사용한다.")

h2("2.6 번들조정 (Bundle Adjustment) — 본 방법")
para("최종 해는 모든 뷰·모든 코너의 재투영 오차를 직접 최소화하는 비선형 최소제곱으로 구한다. 미지수는:")
bullet("X = T_EE_CAM (플랜지←카메라) : 회전 3 + 병진 3")
bullet("T_base_board (베이스 기준 고정 보드 자세) : 회전 3 + 병진 3   → 총 12개 (방식 A)")
para("각 뷰 i에서 보드의 카메라기준 자세를 로봇자세로부터 예측한다:")
mono("T_cam_board(i) = ( T_base_flange(i) · X )⁻¹ · T_base_board")
para("이 자세로 보드 3D코너를 투영해 관측 픽셀과의 잔차를 만들고, 전체 제곱합을 최소화한다:")
mono("min_(X, T_base_board)  Σ_i Σ_c  ρ( ‖ π(K, dist, T_cam_board(i)·P_c) − u_{i,c} ‖² )")
para("K, dist는 방식 A에서 미리 교정해 고정한다. ρ는 robust loss(soft_l1)로 이상치 뷰의 영향을 줄인다. "
     "최적화는 SciPy least_squares(신뢰영역 반사법 'trf')로 수행하며, 회전(rad)과 병진(m)의 스케일 차이를 x_scale로 보정한다. "
     "초기값: intrinsics=calibrateCamera, X=closed-form(Park), 보드자세=코너가 가장 많은 뷰에서 역산.")

h2("2.7 오차의 mm 환산")
para("재투영오차는 픽셀 단위로 나온다. 작은 각도 근사에서 이미지상 1픽셀은 작업거리 Z에서 대략 다음 실거리에 해당한다:")
mono("mm/pixel ≈ Z / fx × 1000        (본 결과: Z≈0.435 m, fx≈888 → ≈0.49 mm/px)")
para("따라서 mm 오차 = (픽셀 오차) × (Z/fx). 뷰마다 Z가 다르므로 뷰별 실제 Z로 환산한다.")

h2("2.8 카메라 → 베이스 → 역기구학(IK) 체인")
para("카메라가 관측한 점 p_camera(예: PnP로 얻은 보드 중심)를 베이스 좌표로 변환:")
mono("p_base = T_base_flange · T_EE_CAM · p_camera")
para("목표 TCP 자세(예: 그 점 위 15 cm, 아래를 향함)를 만들고, 역기구학으로 관절각을 구한다:")
mono("q* = IK(T_base_target)      (RTDE getInverseKinematics, 현재 관절을 초기값으로)")
para("정기구학 FK(q*)가 목표와 일치하면 해가 유효하다. (본 ur-rtde의 getForwardKinematics는 오작동하여, "
     "현재 자세에 대한 IK 라운드트립 IK(현재TCP)=현재관절 로 대체 검증하였다.)")

# ================= 3. 파이프라인 =================
h1("3. 캘리브레이션 파이프라인")
for i,(t,d) in enumerate([
 ("과거 자세 데이터", "보존된 회전·병진 자세별 컬러영상과 TCP 자세를 사용. 현재 capture_calib.py는 receive-only이며 이 보고서는 모션 명령이나 승인 절차가 아님."),
 ("프레이밍", "이미지-자코비안 프로브로 보드를 화면 중앙에 맞추고, 표준거리를 조절해 보드 전체(70코너)가 들어오게 함."),
 ("검출", "모든 뷰에서 ChArUco 코너 검출(11×8, legacy)."),
 ("내부파라미터", "cv2.calibrateCamera로 K, 왜곡 추정 후 고정(방식 A)."),
 ("핸드아이", "번들조정으로 T_EE_CAM과 보드자세를 재투영오차 최소화로 동시 추정."),
 ("결과 저장", "행렬·intrinsics·보드자세·뷰별 오차를 json/txt로 저장."),
 ("데모", "카메라 관측점 → 베이스 → RTDE IK 관절해."),
]):
    bullet(f"단계 {i+1} — {t}: {d}")

# ================= 4. 문제 진단 =================
h1("4. 진행 중 발견·해결한 문제들 (데이터 기반 진단)")
h2("4.1 두 대의 카메라")
para("손목 D435i의 정확한 시리얼은 ignored configs/real.local.yaml에서 명시하며, 시작 시 실제 장치와 일치하지 않으면 실패한다.")
h2("4.2 과거 자세 데이터의 범위")
para("보존된 데이터에는 여러 tilt/translation 자세가 포함되지만, 현재 저장소는 해당 이동을 재현하는 자동 모션 도구를 제공하지 않는다. "
     "새 캡처는 operator-positioned receive-only 절차와 별도의 현장 motion-safety 승인 뒤에만 수행한다.")
h2("4.3 보드 스케일 버그 (첫 보드) — 수학적 진단")
para("같은 자세에서 로봇을 정확히 알려진 거리만큼 병진시키고, 카메라가 관측한 보드 이동량과 비교했다. 순수 병진이면 "
     "‖Δt_camera‖ = ‖Δt_robot‖ 이어야 하는데, 모든 쌍에서 비율이 일정하게 0.712였다:")
mono("로봇 50.0 mm → 카메라 35.5 mm,  100.0 mm → 71.2 mm  (비율 0.712 ≈ 1/√2)")
para("PnP의 병진은 가정한 사각형 크기 SQ_M에 비례한다. 비율 0.712 = 30/42.4 는 실제 사각형이 30 mm가 아니라 ~42.4 mm임을 뜻한다. "
     "A4 종이의 가로:세로 비가 정확히 √2이므로, 세로용 보드를 가로로 채워 인쇄하면 √2배 확대되어 30 mm → 42.4 mm가 된다. "
     "SQ_M을 실측값으로 고치자 병진 뷰 불일치가 사라졌다.")
h2("4.4 직사각형 오해 / 다른 사본")
para("자로 잰 35×42 mm 직사각형은 두 장 중 다른 사본이었다. 실제 촬영에 쓴 보드는 데이터상 정사각(등방)으로, "
     "비등방 모델(35×42)로 맞추면 오차가 오히려 커져 정사각임이 확인됐다.")
h2("4.5 새 보드 — LEGACY 패턴")
para("마커는 검출되는데 ChArUco 코너가 0개였다. OpenCV ≥4.6의 코너-마커 배치 규칙 변경 때문으로, "
     "setLegacyPattern(True)를 주자 70/70 코너가 검출됐다.")
h2("4.6 Protective stop (보호정지)")
para("낮고 뻗은 자세(z≈0.07 m)에서 카메라를 기울이면 tool이 컨베이어/안전평면에 근접해 보호정지가 반복됐다. "
     "홈 관절자세(z≈0.16 m)로 복귀 후 검증된 안전 높이(중심 z≈0.36 m)에서 작업하니 27뷰 전부 정지 없이 촬영됐다. "
     "(보호정지는 고장이 아니라 안전기능이며 하드웨어 손상은 없다.)")

# ================= 5. 방식 A vs B =================
h1("5. 방식 선택: A(고정) vs B(동시최적화)")
para("A: intrinsics를 별도 교정 후 고정 → 미지수 12개. B: intrinsics를 함께 최적화 → 미지수 21개.")
para("B는 자유도가 커 재투영 RMS는 낮아지지만, 이번처럼 뷰가 화면 중앙·비슷한 거리에 몰리면 주점(cx)이 622→574로 "
     "이동하고 그 오차를 카메라 회전(16.7°→20.4°)이 흡수하는 intrinsic–extrinsic 커플링이 생긴다. 수학적으론 잔차 최소지만 "
     "물리적으론 덜 정확하다. 따라서 커플링이 없는 A(별도 intrinsics 고정)를 최종 채택했다 — 정밀 논문의 표준.")

# ================= 6. 최종 결과 =================
h1("6. 최종 결과 (방식 A)")
para(f"사용 뷰 {R['n_views']}개, 코너 {R['n_corners']}개. 내부파라미터(고정):", bold=True)
mono(f"fx={intr['fx']:.2f}  fy={intr['fy']:.2f}  cx={intr['cx']:.2f}  cy={intr['cy']:.2f}\n"
     f"dist(k1,k2,p1,p2,k3)={[round(x,5) for x in intr['dist']]}   (calibrateCamera RMS {intr['rms_px']:.3f} px)")
para("T_EE_CAM (플랜지 좌표계에서 본 카메라 자세; p_flange = T_EE_CAM · p_camera):", bold=True)
mat(T)
para(f"→ 카메라 오프셋 {np.round(np.array(R['camera_offset_mm']),2).tolist()} mm, 플랜지 대비 회전 {R['camera_rot_deg']:.2f}°.")
para("T_base_board (베이스 기준 고정 보드 자세):", bold=True)
mat(Tbb)

h2("6.1 정확도 — 재투영 오차 (px 및 mm)")
tbl = doc.add_table(rows=6, cols=3); tbl.style = "Light Grid Accent 1"
data = [("지표","픽셀","mm"),
 ("전체 RMS", f"{err['reproj_rms_px']:.2f} px", f"{err['reproj_rms_mm']:.2f} mm"),
 ("뷰별 최소(best)", f"{err['per_view_px']['min']:.2f}", f"{err['per_view_mm']['min']:.3f}"),
 ("뷰별 중앙값(median)", f"{err['per_view_px']['median']:.2f}", f"{err['per_view_mm']['median']:.3f}"),
 ("뷰별 평균(mean)", f"{err['per_view_px']['mean']:.2f}", f"{err['per_view_mm']['mean']:.3f}"),
 ("뷰별 최대(worst)", f"{err['per_view_px']['max']:.2f}", f"{err['per_view_mm']['max']:.3f}")]
for rr_,row in enumerate(data):
    for cc,val in enumerate(row):
        cell=tbl.rows[rr_].cells[cc]; cell.text=val
        for pp in cell.paragraphs:
            for run in pp.runs: run.font.size=Pt(10); run.bold=(rr_==0)
para(f"작업거리 {R['working_distance_m']:.3f} m 기준. 스케일·보드·검출 문제는 모두 해결(최선 뷰 {err['per_view_mm']['min']:.3f} mm). "
     "남은 ~2 mm 바닥은 UR5e 절대 정기구학(FK) 오차가 주원인으로 판단된다.")

# ================= 7. 데모 =================
h1("7. 데모 — 카메라 좌표 → 역기구학")
bullet("카메라가 보드 중심 검출 → 카메라계 [15.5, 38.1, 185.0] mm (거리 0.189 m).")
bullet("핸드아이로 베이스계 변환: p_base = [-0.129, -0.590, -0.041] m.")
bullet("목표 TCP(위 15 cm, 아래 향함) → RTDE IK → 관절 [-89.6, -124.6, -106.9, -36.8, 90.0, 0.45]°.")
bullet("검증: 안전한계 내 = True, IK 라운드트립 오차 = 0.000°. → 카메라가 본 물체를 도달 가능한 관절각으로 정확히 변환.")

# ================= 7.1 실증 결과 =================
h1("7.1 과거 실증 결과 — 보관된 측정 데이터")
para("캘리브레이션이 실제로 작동하는지 로봇을 움직여 두 가지로 검증하였다.")

h2("PHASE 1 — 도달 테스트 (지령 목표 vs 실제 도착, 축별 비교)")
para("보드를 한 번 관측해 베이스 좌표계에서의 보드 자세를 구한 뒤, 보드 위 5개 지점(중심+4모서리) 위 20 cm로 "
     "TCP를 이동시키고, 지령한 목표 TCP와 실제 도착 TCP를 축별로 비교하였다.")
t1 = doc.add_table(rows=len(D["phase1_reach"])+1, cols=5); t1.style="Light Grid Accent 1"
hdr=["지점","목표 (x, y, z) m","도착 (x, y, z) m","축별 오차 dx,dy,dz (mm)","|오차| (mm)"]
for c,v in enumerate(hdr):
    cell=t1.rows[0].cells[c]; cell.text=v
    for p in cell.paragraphs:
        for r_ in p.runs: r_.bold=True; r_.font.size=Pt(9)
for ri,row in enumerate(D["phase1_reach"]):
    vals=[row["point"],
          ", ".join(f"{x:.4f}" for x in row["target"]),
          ", ".join(f"{x:.4f}" for x in row["arrived"]),
          ", ".join(f"{x:+.3f}" for x in row["err_axis_mm"]),
          f"{row['err_mm']:.3f}"]
    for c,v in enumerate(vals):
        cell=t1.rows[ri+1].cells[c]; cell.text=v
        for p in cell.paragraphs:
            for r_ in p.runs: r_.font.size=Pt(8.5)
para(f"→ 평균 |오차| {D['phase1_mean_mm']:.3f} mm, 최대 {D['phase1_max_mm']:.3f} mm. 목표와 도착이 마이크로미터 수준으로 일치한다. "
     "이는 로봇이 '카메라가 계산한 목표'에 얼마나 정확히 가는가(로봇 반복정밀도, UR5e ±0.03 mm)를 뜻하며, "
     "그 목표가 실제 물체 위치와 맞는가(핸드아이 정확도)는 아래 PHASE 2에서 별도로 측정한다.", bold=True)

h2("PHASE 2 — 일관성 테스트 (여러 시점에서 재관측한 실제 정확도)")
para("보드 전체가 보이는 5개 시점(top-down + ±8° tilt)에서 각각 보드 원점을 베이스 좌표계로 재추정하였다. "
     "핸드아이가 정확할수록 모든 시점의 추정이 일치한다.")
t2 = doc.add_table(rows=len(D["phase2_consistency"])+1, cols=3); t2.style="Light Grid Accent 1"
hdr=["시점","보드 원점 (베이스, m)","코너 수"]
for c,v in enumerate(hdr):
    cell=t2.rows[0].cells[c]; cell.text=v
    for p in cell.paragraphs:
        for r_ in p.runs: r_.bold=True; r_.font.size=Pt(10)
for ri,row in enumerate(D["phase2_consistency"]):
    vals=[row["view"], str([round(x,4) for x in row["origin"]]), str(row["corners"])]
    for c,v in enumerate(vals):
        cell=t2.rows[ri+1].cells[c]; cell.text=v
        for p in cell.paragraphs:
            for r_ in p.runs: r_.font.size=Pt(10)
para(f"→ 원점 편차(std) {D['phase2_spread_mm']:.2f} mm, 최대편차 {D['phase2_maxdev_mm']:.2f} mm. "
     "어느 로봇 자세에서 봐도 고정된 보드를 ~4 mm 안에서 일치되게 찾는다는 뜻으로, 이것이 실제 동작에서 실증된 "
     "핸드아이 정확도이다 (재투영 RMS 2 mm + 자세별 UR5e 절대 기구학 오차 포함).", bold=True)

# ================= 8. 결론 =================
h1("8. 결론 (Conclusion)")
para("본 프로젝트에서 eye-in-hand 핸드-아이 캘리브레이션을 성공적으로 완료하고, 실제 로봇 동작으로 검증하였다.", bold=True)
bullet(f"핸드아이 확정: 카메라는 플랜지 기준 오프셋 {np.round(np.array(R['camera_offset_mm']),1).tolist()} mm, 회전 {R['camera_rot_deg']:.1f}°에 장착.")
bullet(f"정확도(재투영): RMS {err['reproj_rms_mm']:.2f} mm (최선 뷰 {err['per_view_mm']['min']:.2f} mm), 방식 A(내부파라미터 고정 + 번들조정).")
bullet(f"실동작 검증 ①(로봇 이동 정확도): 지령 목표 vs 실제 도착이 평균 {D['phase1_mean_mm']:.3f} mm(최대 {D['phase1_max_mm']:.3f} mm)로 일치 — 로봇은 계산된 목표에 마이크로미터 수준으로 도달.")
bullet(f"실동작 검증 ②(핸드아이 정확도): 5개 시점에서 보드 위치를 {D['phase2_spread_mm']:.1f} mm(최대 {D['phase2_maxdev_mm']:.1f} mm) 이내로 일치 추정 — 이것이 실제 물체 위치화 정확도.")
bullet("두 정확도는 다른 개념이다: 카메라가 본 물체를 로봇이 실제로 찍는 전체 정확도는 핸드아이(~2–4 mm)가 지배하며, 로봇 이동(~0.015 mm)은 거의 오차를 더하지 않는다.")
bullet("과거 데이터 수집에서 발생한 카메라 선택, 보드 스케일 √2 오류, 잘못된 사본, legacy 패턴, 보호정지를 데이터 기반으로 진단했다.")
bullet("판정: 오프라인 도달성 분석 자료로 보존한다. 비전 기반 실제 로봇 동작은 별도의 최신 캘리브레이션과 deployment evidence 전부를 통과하기 전까지 승인되지 않는다.")
bullet("정밀검사 논문 본값으로 ~2 mm 미만을 노린다면: 전용 intrinsics 세션 + 로봇 기구학(DH) 캘리 + 독립 held-out 검증을 추가할 것.")

# ================= 9. 권장 =================
h1("9. 정밀검사 논문을 위한 권장사항")
bullet("전용 intrinsics 세션: 보드를 화면 전역(가장자리)·여러 거리·여러 기울기로 촬영해 intrinsics를 탄탄히 고정 후 핸드아이 → 커플링 제거.")
bullet("로봇 기구학 캘리: ~2 mm가 남으면 번들조정에 DH 보정항을 추가해 UR5e 절대정확도 오차를 식별.")
bullet("독립 검증: 내부 재투영 RMS 외에, 알려진 점 터치/거리 측정 등 held-out 지표를 함께 보고.")
bullet("계측용 타깃: 평판(유리/알루미늄)에 부착, 사각형 크기를 캘리퍼로 실측해 코드에 정확히 입력.")
bullet("보고 방법: closed-form이 아니라 재투영 기반 번들조정으로 보고.")

# ================= 10. 산출물 =================
h1("10. 산출물")
bullet("calibrate_handeye_final.py — 최종 재현 스크립트(방식 A).")
bullet("capture_calib.py — config-pinned receive-only calibration capture.")
bullet("bundle_adjust.py / refine_and_calibrate.py / deployment/Calib_Verify.py.")
bullet("captures_final/results/ — T_EE_CAM.json·.txt, intrinsics.json, calibration_results.json, summary.txt.")
bullet("calibration_report.pptx (발표용), calibration_report.docx (본 문서).")

doc.save("calibration_report.docx")
print("saved calibration_report.docx")
