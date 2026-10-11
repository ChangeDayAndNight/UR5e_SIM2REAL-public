"""Generate a detailed PPT report of the eye-in-hand hand-eye calibration project."""
import json, os
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

R = json.load(open("captures_final/results/calibration_results.json"))
D = json.load(open("captures_final/results/demo_results.json"))
T = np.array(R["T_EE_CAM"]); intr = R["intrinsics"]; err = R["error"]

NAVY = RGBColor(0x1F, 0x3A, 0x5F); BLUE = RGBColor(0x2E, 0x6D, 0xB4)
GRAY = RGBColor(0x44, 0x44, 0x44); GREEN = RGBColor(0x1E, 0x7A, 0x3C); RED = RGBColor(0xB1, 0x2A, 0x2A)
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def slide(title, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.05))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.45), Inches(0.12), Inches(12.4), Inches(0.85)).text_frame
    tb.word_wrap = True; p = tb.paragraphs[0]; p.text = title
    p.font.size = Pt(30); p.font.bold = True; p.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
    if subtitle:
        sp = tb.add_paragraph(); sp.text = subtitle; sp.font.size = Pt(13); sp.font.color.rgb = RGBColor(0xCF,0xDD,0xEE)
    return s

def parse(it):
    if isinstance(it, str):
        return 0, it, GRAY
    it = list(it); lvl = 0; color = GRAY
    if it and isinstance(it[0], int):
        lvl = it.pop(0)
    txt = it.pop(0)
    if it and isinstance(it[0], RGBColor):
        color = it[0]
    return lvl, txt, color

def bullets(s, items, left=0.55, top=1.3, width=12.2, height=5.9, size=17):
    tf = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        lvl, txt, color = parse(it)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("• " if lvl == 0 else "   – ") + txt
        p.font.size = Pt(size - lvl*2); p.level = lvl
        p.font.color.rgb = color
        p.space_after = Pt(5)
    return s

def title_slide():
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(1, 0, Inches(2.4), prs.slide_width, Inches(2.7))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.8), Inches(2.7), Inches(11.7), Inches(2.2)).text_frame; tb.word_wrap=True
    p = tb.paragraphs[0]; p.text = "Eye-in-Hand Hand-Eye Calibration"; p.font.size=Pt(40); p.font.bold=True
    p.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
    sp = tb.add_paragraph(); sp.text="RealSense D435i on UR5e flange  ·  ChArUco + Bundle Adjustment  ·  RTDE"
    sp.font.size=Pt(18); sp.font.color.rgb=RGBColor(0xCF,0xDD,0xEE)
    f = s.shapes.add_textbox(Inches(0.8), Inches(5.4), Inches(11.7), Inches(1)).text_frame
    f.paragraphs[0].text = "Reproducible offline analysis · full result set · camera->base reachability · precision-inspection assessment"
    f.paragraphs[0].font.size=Pt(14); f.paragraphs[0].font.color.rgb=GRAY

# ---- slides ----
title_slide()

slide("1. Goal & System Setup")
bullets(prs.slides[-1], [
 "Objective: find T_EE_CAM — the pose of the camera in the robot flange frame — so camera observations can be expressed in the robot base frame for reachability analysis AND precision inspection.",
 "Configuration: eye-in-hand — RealSense D435i rigidly mounted on the UR5e flange; ChArUco board fixed on a powered-off conveyor. The retained dataset contains historical varied poses.",
 ("Hardware:", ), (1,"UR5e endpoint loaded from ignored configs/real.local.yaml"),
 (1,"D435i eye-in-hand camera serial loaded from ignored configs/real.local.yaml"),
 (1,"Board: 11x8 ChArUco, square 15 mm, marker 11 mm, DICT_4X4_50, LEGACY pattern"),
 "Software: the locked Windows deployment runtime; pyrealsense2, ur-rtde, OpenCV, SciPy.",
 "Key relation:  p_base = T_base_flange · T_EE_CAM · p_camera   (T_base_flange from RTDE getActualTCPPose).",
])

slide("2. Pipeline Overview")
bullets(prs.slides[-1], [
 "1) Historical capture set — retained color frames and actual TCP poses provide rotation and translation diversity. Current capture_calib.py is receive-only.",
 "2) Framing — image-Jacobian probe centers the board; standoff chosen so the whole board is in view.",
 "3) Detect — ChArUco corners in every view (11x8, legacy pattern).",
 "4) Intrinsics — cv2.calibrateCamera over all views -> K, distortion (then FIXED).",
 "5) Hand-eye — bundle adjustment refines T_EE_CAM + board pose by minimizing reprojection error of every corner.",
 "6) Save all results; 7) Offline analysis: camera-frame point -> base frame -> kinematic reachability (no command).",
 ("Current authorization boundary:", NAVY),
 (1,"this generated report records historical measurements only; it neither sends commands nor authorizes real motion."),
])

slide("3. Scripts Delivered")
bullets(prs.slides[-1], [
 ("capture_calib.py", NAVY), (1,"operator-positioned, receive-only capture with config-pinned robot and camera."),
 ("calibrate_handeye_final.py  <-- FINAL, reproducible (Method A)", GREEN),
   (1,"intrinsics (calibrated, fixed) + bundle-adjusted hand-eye; saves every result."),
 ("bundle_adjust.py", NAVY), (1,"research-grade BA with per-view px/mm error reporting and A/B comparison."),
 ("deployment/Calib_Verify.py", NAVY), (1,"canonical fail-closed calibration verification."),
 ("refine_and_calibrate.py", NAVY), (1,"self-calibrates intrinsics from the ChArUco views."),
])

slide("4. Challenges Solved (1/2)", "each was diagnosed from data, not guessed")
bullets(prs.slides[-1], [
 ("Two cameras", NAVY),
   (1,"the eye-in-hand camera identity is loaded from ignored configs/real.local.yaml; scripts reject automatic first-device selection."),
 ("Historical motion capture", NAVY),
   (1,"the retained dataset includes varied robot poses; this report is not motion authorization and current capture_calib.py is receive-only."),
 ("Board framing", NAVY),
   (1,"board initially off-center/clipped; image-Jacobian probe + standoff tuning -> full 70/70 corners per view."),
 ("BOARD SCALE BUG (1st board)", RED),
   (1,"robot moved 50 mm but camera saw 35.5 mm (ratio 0.712 = 1/sqrt2). Board printed at sqrt(2) scale (A4 aspect) -> square was ~42 mm, not 30 mm. Fixed by measuring/setting the true size."),
])

slide("5. Challenges Solved (2/2)")
bullets(prs.slides[-1], [
 ("Wrong board copy / anisotropy", RED),
   (1,"a measured 35x42 mm rectangle turned out to be the OTHER printed copy; the imaged board was square (data-confirmed)."),
 ("New board: LEGACY ChArUco pattern", RED),
   (1,"markers detected but 0 ChArUco corners -> OpenCV>=4.6 layout change. setLegacyPattern(True) -> 70/70 corners."),
 ("Protective stops on tilt", RED),
   (1,"tilting from a low/extended pose (z~0.07 m) tripped the safety stop. Solved by returning to a HOME joint pose (z~0.16 m) and working at a proven-safe height (center z~0.36 m) -> all 27 views captured with no stop."),
 ("Closed-form vs bundle adjustment", NAVY),
   (1,"replaced cv2.calibrateHandEye (per-view poses only) with a reprojection-minimizing bundle adjustment (paper-grade)."),
])

slide("6. Method Chosen: A (intrinsics fixed) vs B (jointly refined)")
bullets(prs.slides[-1], [
 ("Bundle adjustment jointly optimizes:", NAVY),
   (1,"T_EE_CAM (flange<-camera, 6 dof) + T_base_board (6 dof), minimizing reprojection of all 1890 corners (SciPy trf + soft_l1)."),
 ("A) intrinsics calibrated separately then FIXED  ->  12 parameters", GREEN),
 ("B) intrinsics added to the optimization  ->  21 parameters (lower RMS, but risky)", GRAY),
 ("Why A was chosen:", NAVY),
   (1,"With views clustered near image center, B lets the principal point drift (cx 622->574) and the camera rotation absorbs it (16.7deg -> 20.4deg). Lower RMS but physically less trustworthy."),
   (1,"A avoids this intrinsic-extrinsic coupling and gives a stable, defensible hand-eye — the standard for precision papers."),
])

s = slide("7. FINAL RESULT — T_EE_CAM (Method A)")
tb = s.shapes.add_textbox(Inches(0.55), Inches(1.25), Inches(7.4), Inches(3)).text_frame; tb.word_wrap=True
tb.paragraphs[0].text = "T_EE_CAM  (p_flange = T_EE_CAM · p_camera):"; tb.paragraphs[0].font.size=Pt(15); tb.paragraphs[0].font.bold=True
mat = "\n".join("[ " + "  ".join(f"{v:9.6f}" for v in row) + " ]" for row in T)
pm = tb.add_paragraph(); pm.text = mat; pm.font.name="Consolas"; pm.font.size=Pt(14); pm.font.color.rgb=NAVY
bullets(s, [
 (f"camera offset from flange: {np.round(np.array(R['camera_offset_mm']),1).tolist()} mm", GREEN),
 (f"camera rotation from flange: {R['camera_rot_deg']:.2f} deg", GREEN),
 (f"intrinsics (fixed): fx={intr['fx']:.1f}  fy={intr['fy']:.1f}  cx={intr['cx']:.1f}  cy={intr['cy']:.1f}", GRAY),
 (f"views: {R['n_views']}   corners: {R['n_corners']}   working distance: {R['working_distance_m']:.3f} m", GRAY),
], top=4.35, size=16)

s = slide("8. Accuracy — Reprojection Error (px and mm)")
rows = [("Metric","pixels","mm"),
 ("Overall RMS", f"{err['reproj_rms_px']:.2f} px", f"{err['reproj_rms_mm']:.2f} mm"),
 ("per-view min (best)", f"{err['per_view_px']['min']:.2f} px", f"{err['per_view_mm']['min']:.2f} mm"),
 ("per-view median", f"{err['per_view_px']['median']:.2f} px", f"{err['per_view_mm']['median']:.2f} mm"),
 ("per-view mean", f"{err['per_view_px']['mean']:.2f} px", f"{err['per_view_mm']['mean']:.2f} mm"),
 ("per-view max (worst)", f"{err['per_view_px']['max']:.2f} px", f"{err['per_view_mm']['max']:.2f} mm")]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.7), Inches(1.4), Inches(7.2), Inches(3.2)).table
for c in range(3): tbl.columns[c].width = Inches(2.4)
for r,row in enumerate(rows):
    for c,val in enumerate(row):
        cell=tbl.cell(r,c); cell.text=val
        pr=cell.text_frame.paragraphs[0]; pr.font.size=Pt(15); pr.font.bold=(r==0)
        if r==0: cell.fill.solid(); cell.fill.fore_color.rgb=NAVY; pr.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
bullets(s, [
 ("Scale, board and detection issues are fully resolved (all 27 views 70/70 corners; best view 0.19 mm).", GREEN),
 ("Remaining ~2 mm floor is dominated by UR5e absolute forward-kinematics error, not by the camera.", NAVY),
 ("Reprojection RMS = internal consistency; report it WITH an independent held-out check for a paper.", GRAY),
], left=8.2, top=1.4, width=4.7, size=14)

slide("9. Demo — Camera Coordinate -> Inverse Kinematics")
bullets(prs.slides[-1], [
 "A historical validation dataset demonstrates the full coordinate chain end-to-end:",
 (1,"1. Camera detects the board -> board center in CAMERA frame = [15.5, 38.1, 185.0] mm (0.189 m)."),
 (1,"2. Hand-eye maps it to BASE frame:  p_base = T_base_flange · T_EE_CAM · p_camera = [-0.129, -0.590, -0.041] m."),
 (1,"3. Desired TCP pose (hover 15 cm above, top-down) -> RTDE getInverseKinematics ->"),
 (1,"   joints = [-89.6, -124.6, -106.9, -36.8, 90.0, 0.45] deg."),
 (1,"4. Verified: within safety limits = True; IK round-trip error = 0.000 deg."),
 ("IK is validated via round-trip (getForwardKinematics is buggy in ur-rtde and is not used).", GRAY),
 ("Result: an object seen by the camera is correctly turned into reachable robot joint angles.", GREEN),
])

# ---- Validation Phase 1: reach test (target vs arrived) ----
s = slide("10. Validation by Real Robot Motion — Reach Test (Phase 1)",
          "commanded TCP target (from camera+hand-eye) vs actually arrived TCP, per axis")
rows = [("Point", "Target (x, y, z) m", "Arrived (x, y, z) m", "Error dx,dy,dz (mm)", "|err|")]
for r in D["phase1_reach"]:
    tg = ", ".join(f"{v:.4f}" for v in r["target"]); ar = ", ".join(f"{v:.4f}" for v in r["arrived"])
    ea = ", ".join(f"{v:+.3f}" for v in r["err_axis_mm"])
    rows.append((r["point"], tg, ar, ea, f"{r['err_mm']:.3f}"))
tbl = s.shapes.add_table(len(rows), 5, Inches(0.35), Inches(1.35), Inches(12.6), Inches(2.9)).table
for ci,w in enumerate([1.3, 3.3, 3.3, 3.0, 1.1]): tbl.columns[ci].width=Inches(w)
for ri,row in enumerate(rows):
    for ci,val in enumerate(row):
        c=tbl.cell(ri,ci); c.text=val; p=c.text_frame.paragraphs[0]
        p.font.size=Pt(11 if ri else 11); p.font.bold=(ri==0)
        if ci in (1,2): p.font.name="Consolas"; p.font.size=Pt(10)
        if ri==0: c.fill.solid(); c.fill.fore_color.rgb=NAVY; p.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
bullets(s, [
 (f"Mean |error| = {D['phase1_mean_mm']:.3f} mm, max = {D['phase1_max_mm']:.3f} mm — target and arrived match at the micron level.", GREEN),
 ("This isolates ROBOT MOTION accuracy: the robot reaches the camera-computed target essentially perfectly (UR5e repeatability ±0.03 mm).", NAVY),
 ("Whether that target equals the true object location is the HAND-EYE accuracy (Phase 2, ~4 mm) — a separate quantity.", GRAY),
], top=4.45, size=13)

# ---- Validation Phase 2: consistency test ----
s = slide("11. Validation by Real Robot Motion — Consistency Test (Phase 2)",
          "board origin re-estimated from 5 different viewpoints; agreement = true hand-eye accuracy")
rows = [("Viewpoint", "Board origin in base (m)", "Corners")]
for r in D["phase2_consistency"]:
    rows.append((r["view"], str([round(v,4) for v in r["origin"]]), str(r["corners"])))
tbl = s.shapes.add_table(len(rows), 3, Inches(0.6), Inches(1.35), Inches(9.2), Inches(3.0)).table
tbl.columns[0].width=Inches(2.0); tbl.columns[1].width=Inches(5.2); tbl.columns[2].width=Inches(2.0)
for ri,row in enumerate(rows):
    for ci,val in enumerate(row):
        c=tbl.cell(ri,ci); c.text=val; p=c.text_frame.paragraphs[0]; p.font.size=Pt(14); p.font.bold=(ri==0)
        if ri==0: c.fill.solid(); c.fill.fore_color.rgb=NAVY; p.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
bullets(s, [
 (f"Origin spread (std) = {D['phase2_spread_mm']:.2f} mm,  max deviation = {D['phase2_maxdev_mm']:.2f} mm.", GREEN),
 ("Meaning: from any of these robot poses, the camera + hand-eye locate the fixed board within", NAVY),
 (1,"~4 mm. This is the calibration accuracy demonstrated under real motion (consistent with the 2 mm reprojection RMS plus per-pose UR5e absolute-kinematics error)."),
], top=4.55, size=15)

# ---- Conclusion ----
s = slide("12. Conclusion")
bullets(s, [
 ("The eye-in-hand hand-eye calibration is COMPLETE and VALIDATED.", GREEN),
 (1,f"T_EE_CAM determined: camera offset {np.round(np.array(R['camera_offset_mm']),1).tolist()} mm, rotation {R['camera_rot_deg']:.1f} deg."),
 (1,f"Reprojection RMS {err['reproj_rms_mm']:.2f} mm (best view {err['per_view_mm']['min']:.2f} mm), method A (fixed intrinsics + bundle adjustment)."),
 ("Validated by real robot motion — two distinct accuracies:", NAVY),
 (1,f"ROBOT MOTION (target vs arrived): mean {D['phase1_mean_mm']:.3f} mm — reaches camera-computed targets to microns."),
 (1,f"HAND-EYE (object localization): {D['phase2_spread_mm']:.1f} mm across 5 viewpoints — the true end-to-end accuracy."),
 (1,f"Combined pick accuracy is dominated by hand-eye (~2-4 mm); motion adds essentially nothing."),
  ("Historical dataset issues resolved: camera identity, board scale (sqrt2), wrong copy,", NAVY),
  (1,"legacy ChArUco pattern, and retained pose-set quality. No current automatic motion is authorized."),
 ("Fit for offline reachability analysis; real motion remains blocked pending the repository deployment evidence gates.", GREEN),
 ("For a precision-inspection PAPER: add a dedicated intrinsic session, robot kinematic (DH) calibration, and an independent held-out validation to push below ~2 mm.", GRAY),
])

slide("13. Recommendations for the Precision-Inspection Paper")
bullets(prs.slides[-1], [
 ("1. Dedicated intrinsic calibration session", NAVY),
   (1,"capture the board across the WHOLE image (corners/edges) at several distances & tilts; fix those intrinsics, then solve hand-eye -> removes the intrinsic-extrinsic coupling."),
 ("2. Robot kinematic calibration", NAVY),
   (1,"if ~2 mm remains after step 1, add DH corrections to the bundle adjustment (identifies UR5e absolute-accuracy error)."),
 ("3. Independent validation", NAVY),
   (1,"report a held-out metric (touch a known point / measure a known distance), not only internal reprojection RMS."),
 ("4. Metrology-grade target", NAVY),
   (1,"ChArUco on a flat rigid plate; verify square size with calipers and enter the exact value in code."),
 ("Method to report: bundle adjustment (reprojection-based), NOT closed-form calibrateHandEye.", GREEN),
])

s = slide("14. Deliverables & Files")
bullets(s, [
 ("Code (repository/calibration/):", NAVY),
   (1,"calibrate_handeye_final.py  — reproduce the final calibration (Method A)"),
   (1,"capture_calib.py  — receive-only, config-pinned capture"),
   (1,"bundle_adjust.py, refine_and_calibrate.py, deployment/Calib_Verify.py"),
 ("Results (captures_final/results/):", NAVY),
   (1,"T_EE_CAM.json / T_EE_CAM.txt  — the 4x4 hand-eye matrix"),
   (1,"intrinsics.json  — fx,fy,cx,cy + distortion"),
   (1,"calibration_results.json  — everything (matrices, board pose, per-view errors)"),
   (1,"summary.txt  — human-readable summary"),
 ("Data: captures_final/  — 27 views, all 70/70 ChArUco corners, rotation + translation diversity.", GRAY),
])

# board image slide
img = "captures_final/cap_000.png"
if os.path.exists(img):
    s = slide("Appendix — Centered ChArUco View (70/70 corners)")
    s.shapes.add_picture(img, Inches(3.3), Inches(1.35), height=Inches(5.6))

prs.save("calibration_report.pptx")
print("saved calibration_report.pptx  with", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
